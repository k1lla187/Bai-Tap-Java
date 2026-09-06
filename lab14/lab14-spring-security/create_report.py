from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

ROOT = Path(r"D:\Bai-Tap-Java\lab14\lab14-spring-security")
OUTPUT = ROOT / "Bao-cao-Lab14-Spring-Security.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    r.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.35)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(8.5)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p._p.get_or_add_pPr().append(shd)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)

def add_image_placeholder(doc, title, instruction):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    shade(cell, "F6F8FA")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    r = p.add_run("[VỊ TRÍ CHÈN ẢNH MINH CHỨNG]\n")
    r.bold = True
    r.font.color.rgb = RGBColor(20, 93, 160)
    r.font.size = Pt(12)
    r2 = p.add_run(title + "\n" + instruction)
    r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(35)
    add_caption(doc, f"Hình: {title}")

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(12)
styles["Normal"].paragraph_format.line_spacing = 1.35
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in [("Title", 20, (20, 93, 160)), ("Heading 1", 15, (20, 93, 160)), ("Heading 2", 13, (217, 119, 6))]:
    style = styles[name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(*color)

# Cover page
for text, size, bold in [
    ("TRƯỜNG ĐẠI HỌC ĐẠI NAM", 15, True),
    ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True),
    ("", 12, False),
    ("BÁO CÁO THỰC HÀNH", 19, True),
    ("HỌC PHẦN CÔNG NGHỆ JAVA", 17, True),
    ("", 12, False),
    ("LAB 14: BẢO MẬT ỨNG DỤNG VỚI SPRING SECURITY", 18, True),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if "LAB 14" in text:
        r.font.color.rgb = RGBColor(20, 93, 160)

doc.add_paragraph("\n\n")
info = doc.add_table(rows=4, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.style = "Table Grid"
for row, (label, value) in zip(info.rows, [
    ("Sinh viên thực hiện", "Vi Anh Tuan"),
    ("Mã sinh viên", "20230199"),
    ("Lớp", "DCCNTT14.10"),
    ("Năm học", "2026"),
]):
    set_cell_text(row.cells[0], label, bold=True)
    set_cell_text(row.cells[1], value)
    shade(row.cells[0], "E8F1FB")
doc.add_paragraph("\n\n\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Hà Nội, tháng 09 năm 2026")
r.italic = True

doc.add_page_break()

# Contents
p = doc.add_paragraph("MỤC LỤC", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for line in [
    "1. Giới thiệu bài lab", "2. Mục tiêu", "3. Công nghệ và môi trường", "4. Thiết kế ứng dụng",
    "5. Cài đặt các chức năng", "6. Kiểm thử và kết quả", "7. Kết luận", "Phụ lục: Hướng dẫn chạy"
]:
    doc.add_paragraph(line, style="List Bullet")
doc.add_page_break()

# Sections

doc.add_heading("1. Giới thiệu bài lab", level=1)
doc.add_paragraph("Lab 14 thuộc Chương 4: Phát triển ứng dụng với Spring Framework. Bài thực hành xây dựng một ứng dụng Spring Boot quản lý sinh viên và môn học, trong đó Spring Security được dùng để xác thực người dùng khi đăng nhập và phân quyền truy cập theo vai trò.")
doc.add_paragraph("Ứng dụng có hai vai trò: ADMIN được phép quản lý sinh viên và môn học; USER chỉ được xem và tìm kiếm sinh viên. Giao diện Thymeleaf thay đổi theo quyền của tài khoản đăng nhập.")

doc.add_heading("2. Mục tiêu", level=1)
for item in [
    "Tích hợp Spring Security vào ứng dụng Spring Boot.",
    "Xây dựng trang đăng nhập tùy chỉnh và chức năng đăng xuất.",
    "Cấu hình người dùng trong bộ nhớ với mật khẩu được mã hóa BCrypt.",
    "Phân quyền URL theo vai trò ADMIN và USER.",
    "Bảo vệ chức năng thêm, sửa và xóa dữ liệu.",
    "Hiển thị tên người dùng đang đăng nhập và ẩn/hiện menu theo role.",
    "Tạo trang báo lỗi 403 cho trường hợp không đủ quyền truy cập.",
]: doc.add_paragraph(item, style="List Bullet")

doc.add_heading("3. Công nghệ và môi trường", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(table.rows[0].cells, ["Công nghệ", "Mục đích sử dụng"]):
    set_cell_text(cell, text, True, (255,255,255)); shade(cell, "145DA0")
for left, right in [
    ("Java 17", "Biên dịch và chạy ứng dụng"),
    ("Spring Boot 3.2.0", "Khung phát triển ứng dụng web"),
    ("Spring Security 6", "Xác thực và phân quyền"),
    ("Thymeleaf + Thymeleaf Extras", "Tạo giao diện và kiểm tra role trong view"),
    ("Spring Data JPA + H2", "Lưu trữ dữ liệu mẫu trong cơ sở dữ liệu in-memory"),
    ("Maven", "Quản lý dependency, build và chạy project"),
]:
    cells = table.add_row().cells; set_cell_text(cells[0], left); set_cell_text(cells[1], right)

doc.add_heading("4. Thiết kế ứng dụng", level=1)
doc.add_heading("4.1. Cấu trúc project", level=2)
add_code(doc, "lab14-spring-security/\n├── pom.xml\n├── src/main/java/vn/edu/eaut/lab14/\n│   ├── config/SecurityConfig.java\n│   ├── controller/{Auth,Home,Student,Course}Controller.java\n│   ├── entity/{Student,Course}.java\n│   └── repository/{Student,Course}Repository.java\n└── src/main/resources/templates/\n    ├── auth/login.html\n    ├── error/403.html\n    ├── students/{list,form}.html\n    └── courses/{list,form}.html")

doc.add_heading("4.2. Authentication - Xác thực người dùng", level=2)
doc.add_paragraph("Authentication là quá trình kiểm tra danh tính người dùng. Trong bài lab, hai tài khoản được tạo bằng InMemoryUserDetailsManager. Spring Security nhận username và password từ form POST /login, sau đó đối chiếu mật khẩu với giá trị BCrypt đã mã hóa.")
add_code(doc, "admin / 123456  → role ADMIN\nuser  / 123456  → role USER")
doc.add_paragraph("Mật khẩu được mã hóa bằng BCryptPasswordEncoder, không được sử dụng ở dạng văn bản thuần trong UserDetails.")

doc.add_heading("4.3. Authorization - Phân quyền truy cập", level=2)
doc.add_paragraph("Authorization là quá trình kiểm tra quyền sau khi người dùng đã đăng nhập. SecurityFilterChain xác định role cần thiết cho từng nhóm URL.")
table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(table.rows[0].cells, ["URL", "Quyền truy cập"]): set_cell_text(cell, text, True, (255,255,255)); shade(cell, "145DA0")
for left, right in [
    ("/ , /about , /login , /css/**", "Công khai"),
    ("/students", "ADMIN hoặc USER"),
    ("/students/create, /students/edit/**, /students/save, /students/delete/**", "Chỉ ADMIN"),
    ("/courses/**", "Chỉ ADMIN"),
    ("Các URL còn lại", "Yêu cầu đăng nhập"),
]:
    cells = table.add_row().cells; set_cell_text(cells[0], left); set_cell_text(cells[1], right)

doc.add_heading("4.4. Cấu hình Spring Security", level=2)
add_code(doc, ".requestMatchers(\"/courses/**\").hasRole(\"ADMIN\")\n.requestMatchers(\"/students/create\", \"/students/edit/**\",\n                 \"/students/save\", \"/students/delete/**\")\n    .hasRole(\"ADMIN\")\n.requestMatchers(\"/students/**\").hasAnyRole(\"ADMIN\", \"USER\")\n.formLogin(form -> form.loginPage(\"/login\")\n    .defaultSuccessUrl(\"/students\", true))\n.exceptionHandling(exception -> exception.accessDeniedPage(\"/error/403\"))")

doc.add_heading("4.5. Bảo vệ giao diện bằng Thymeleaf Security", level=2)
doc.add_paragraph("Các template sử dụng namespace sec của thư viện thymeleaf-extras-springsecurity6. Việc ẩn nút cải thiện trải nghiệm người dùng; việc bảo vệ thực sự vẫn do Spring Security thực hiện ở tầng URL.")
add_code(doc, "<a sec:authorize=\"hasRole('ADMIN')\" th:href=\"@{/courses}\">\n    Môn học (ADMIN)\n</a>\n<span sec:authentication=\"name\"></span>\n<form sec:authorize=\"isAuthenticated()\" th:action=\"@{/logout}\" method=\"post\">...</form>")

doc.add_heading("5. Cài đặt các chức năng", level=1)
for title, body in [
    ("Đăng nhập và đăng xuất", "AuthController trả về template auth/login. Spring Security tiếp nhận POST /login; khi thành công, ứng dụng chuyển đến /students. Form POST /logout kết thúc phiên người dùng và quay về trang chủ."),
    ("Quản lý sinh viên", "StudentController cho phép cả ADMIN và USER xem, tìm kiếm danh sách sinh viên. Các endpoint create, edit, save và delete được SecurityConfig giới hạn cho ADMIN. Thao tác xóa dùng POST, nhờ đó giữ CSRF protection mặc định."),
    ("Quản lý môn học", "CourseController nằm dưới /courses/**. Toàn bộ khu vực này dành riêng cho ADMIN, đáp ứng bài tập mở rộng số 6."),
    ("Xử lý lỗi 403", "Khi USER cố truy cập URL của ADMIN, AccessDeniedHandler chuyển đến /error/403 để hiển thị thông báo rõ ràng thay vì trang lỗi mặc định."),
]:
    doc.add_heading(title, level=2); doc.add_paragraph(body)

doc.add_heading("6. Kiểm thử và kết quả", level=1)
doc.add_paragraph("Ứng dụng được build thành công bằng lệnh mvn clean package. Các luồng quan trọng được kiểm thử thủ công trên trình duyệt.")
table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(table.rows[0].cells, ["STT", "Kịch bản", "Kết quả mong đợi", "Kết quả"]): set_cell_text(cell, text, True, (255,255,255)); shade(cell, "145DA0")
for row in [
    ("1", "Đăng nhập admin / 123456", "Vào /students và thấy menu, nút quản lý", "Đạt"),
    ("2", "Đăng nhập user / 123456", "Chỉ xem/tìm kiếm sinh viên", "Đạt"),
    ("3", "USER mở /courses", "Chuyển đến trang 403", "Đạt"),
    ("4", "ADMIN tạo hoặc xóa sinh viên", "Thao tác thành công", "Đạt"),
    ("5", "mvn clean package", "Build thành công", "Đạt"),
]:
    cells = table.add_row().cells
    for cell, value in zip(cells, row): set_cell_text(cell, value)

add_image_placeholder(doc, "Đăng nhập bằng tài khoản ADMIN", "Chèn ảnh chụp trang /students sau khi đăng nhập admin / 123456; ảnh cần thấy menu Môn học và các nút Thêm, Sửa, Xóa.")
add_image_placeholder(doc, "Đăng nhập bằng tài khoản USER", "Chèn ảnh chụp trang /students sau khi đăng nhập user / 123456; ảnh cần cho thấy menu Môn học và các nút quản lý đã bị ẩn.")
add_image_placeholder(doc, "USER bị chặn khi truy cập chức năng ADMIN", "Chèn ảnh trang 403 khi user truy cập http://localhost:8080/courses.")

doc.add_heading("7. Kết luận", level=1)
doc.add_paragraph("Lab 14 đã hoàn thành các yêu cầu về Spring Security: đăng nhập, đăng xuất, người dùng trong bộ nhớ, mã hóa mật khẩu BCrypt, phân quyền ADMIN/USER, bảo vệ URL, ẩn/hiện chức năng bằng Thymeleaf Security và xử lý lỗi 403. Ứng dụng cũng có dữ liệu H2 mẫu để kiểm thử nhanh.")
doc.add_paragraph("Hướng phát triển: thay InMemoryUserDetailsManager bằng UserDetailsService lấy người dùng, mật khẩu mã hóa và role từ cơ sở dữ liệu MySQL để thực hiện bài tập mở rộng số 10.")

doc.add_heading("Phụ lục: Hướng dẫn chạy", level=1)
doc.add_paragraph("Mở PowerShell tại thư mục project và thực hiện các lệnh sau:")
add_code(doc, "cd D:\\Bai-Tap-Java\\lab14\\lab14-spring-security\nmvn clean package\nmvn spring-boot:run")
doc.add_paragraph("Sau khi ứng dụng khởi động, mở trình duyệt tại http://localhost:8080. Tài khoản kiểm thử: admin / 123456 và user / 123456.")

# Footer
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Lab 14 - Bảo mật ứng dụng với Spring Security | Công nghệ Java")
    footer.runs[0].font.size = Pt(9)

doc.save(OUTPUT)
print(OUTPUT)
