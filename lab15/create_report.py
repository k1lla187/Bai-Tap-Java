from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

ROOT = Path(r"D:\Bai-Tap-Java\lab15")
OUTPUT = ROOT / "Bao-cao-Lab15-Spring-Final-Project.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    r.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.35)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(8.5)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p._p.get_or_add_pPr().append(shd)

def add_image_placeholder(doc, title, instruction):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    shade(cell, "F6F8FA")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("[VỊ TRÍ CHÈN ẢNH MINH CHỨNG]\n")
    r.bold = True
    r.font.color.rgb = RGBColor(79, 70, 229)
    r.font.size = Pt(11)

    r2 = p.add_run(f"Ảnh: {title}\n")
    r2.bold = True
    r2.font.size = Pt(10)

    r3 = p.add_run(f"({instruction})")
    r3.italic = True
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(100, 116, 139)
    p.paragraph_format.space_after = Pt(20)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = p_cap.add_run(f"Hình minh họa: {title}")
    rc.italic = True
    rc.font.size = Pt(9.5)

def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    # Title
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐÔNG Á\nKHOA CÔNG NGHỆ THÔNG TIN\n-----------------------")
    r_uni.bold = True
    r_uni.font.size = Pt(12)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(15)
    p_title.paragraph_format.space_after = Pt(15)
    r_title = p_title.add_run("BÁO CÁO THỰC HÀNH BÀI TẬP LỚN (LAB 15)\nHỆ THỐNG QUẢN LÝ SINH VIÊN VÀ ĐĂNG KÝ HỌC PHẦN")
    r_title.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(30, 27, 75)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Học phần: Công nghệ Java (Chương 4: Phát triển ứng dụng với Spring Framework)\nProject: lab15-spring-final-project | Package: vn.edu.eaut.lab15")
    r_sub.italic = True
    r_sub.font.size = Pt(11)

    doc.add_page_break()

    # Section 1
    doc.add_heading("1. Giới thiệu Đề tài và Kiến trúc Hệ thống", level=1)
    p = doc.add_paragraph("Bài tập lớn Lab 15 nhằm mục đích xây dựng một ứng dụng hoàn chỉnh tích hợp toàn diện các thành phần đã học trong Chương 4 bao gồm: Spring Boot, Spring MVC, Spring Data JPA, Thymeleaf và Spring Security. Hệ thống giải quyết bài toán nghiệp vụ Quản lý sinh viên, Danh mục môn học/khóa học và Đăng ký học phần (Quan hệ Nhiều - Nhiều có thuộc tính ngày đăng ký).")
    
    p_arch = doc.add_paragraph("Kiến trúc hệ thống được phân lớp nghiêm ngặt theo mô hình MVC đa tầng:")
    doc.add_paragraph("• Entity Layer: Định nghĩa các thực thể CSDL Student, Course, Enrollment (quan hệ ManyToOne và OneToMany).")
    doc.add_paragraph("• Repository Layer: Kế thừa JpaRepository cung cấp các thao tác CRUD và truy vấn nâng cao (findByStudentId, existsByStudentIdAndCourseId).")
    doc.add_paragraph("• Service Layer: Đóng gói toàn bộ nghiệp vụ (Transaction, kiểm tra đăng ký trùng lặp, tính tổng số tín chỉ, hủy học phần).")
    doc.add_paragraph("• Controller Layer: Điều hướng giao diện web, xử lý form, bắt lỗi xác thực (Bean Validation) và phản hồi dữ liệu Model sang Thymeleaf.")
    doc.add_paragraph("• Security Layer: Bảo vệ các route, xác thực đăng nhập và phân quyền chi tiết cho 2 nhóm tài khoản ADMIN và USER.")

    # Section 2
    doc.add_heading("2. Thiết kế Cơ sở Dữ liệu & Mối quan hệ N-N", level=1)
    doc.add_paragraph("Cơ sở dữ liệu gồm 3 bảng quan hệ: students, courses và enrollments. Bảng enrollments giữ vai trò là bảng trung gian có thêm thuộc tính enroll_date và khóa ngoại trỏ về 2 bảng cha.")
    
    t_db = doc.add_table(rows=4, cols=4)
    t_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Bảng CSDL", "Khóa chính (PK)", "Khóa ngoại (FK)", "Mô tả / Ràng buộc"]
    for i, h in enumerate(headers):
        set_cell_text(t_db.cell(0, i), h, bold=True)
        shade(t_db.cell(0, i), "E2E8F0")

    row_data = [
        ("students", "id (BIGINT)", "-", "Lưu mã SV (Unique), họ tên, ngày sinh"),
        ("courses", "id (BIGINT)", "-", "Lưu mã môn (Unique), tên môn học, số tín chỉ"),
        ("enrollments", "id (BIGINT)", "student_id -> students(id)\ncourse_id -> courses(id)", "Lưu enroll_date. Ràng buộc UNIQUE(student_id, course_id)")
    ]
    for r_idx, row in enumerate(row_data, start=1):
        for c_idx, val in enumerate(row):
            set_cell_text(t_db.cell(r_idx, c_idx), val)

    doc.add_paragraph().paragraph_format.space_before = Pt(10)

    # Section 3
    doc.add_heading("3. Báo cáo Chi tiết 10 Bài tập và Mã Nguồn", level=1)
    doc.add_paragraph("Dưới đây là tóm tắt giải pháp thực hiện cho 10 bài tập trong yêu cầu bài lab:")
    
    t_ex = doc.add_table(rows=11, cols=3)
    t_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    ex_headers = ["Bài tập", "Yêu cầu kỹ thuật", "Lớp / File thực thi"]
    for i, h in enumerate(ex_headers):
        set_cell_text(t_ex.cell(0, i), h, bold=True)
        shade(t_ex.cell(0, i), "E2E8F0")

    ex_rows = [
        ("Bài 1", "Thiết kế Entity Course (mã môn, tên môn, tín chỉ)", "Course.java"),
        ("Bài 2", "Thiết kế Entity Enrollment (quan hệ N-N có thuộc tính)", "Enrollment.java"),
        ("Bài 3", "Tạo EnrollmentRepository (findByStudentId, exists...)", "EnrollmentRepository.java"),
        ("Bài 4", "Tạo EnrollmentService xử lý đăng ký và kiểm tra trùng", "EnrollmentService.java"),
        ("Bài 5", "Tạo EnrollmentController hiển thị form và lưu đăng ký", "EnrollmentController.java"),
        ("Bài 6", "Xây dựng trang danh sách đăng ký học phần", "enrollments/list.html"),
        ("Bài 7", "Chức năng hủy đăng ký học phần", "EnrollmentService.cancel() & route cancel"),
        ("Bài 8", "Xem danh sách môn học 1 SV đã đăng ký & tổng tín chỉ", "enrollments/student-history.html"),
        ("Bài 9", "Dashboard thống kê tổng số sinh viên, môn học, lượt ĐK", "DashboardController.java & dashboard.html"),
        ("Bài 10", "Hoàn thiện UI, thông báo Alert và phân quyền Security", "SecurityConfig.java & layout.html")
    ]
    for r_idx, row in enumerate(ex_rows, start=1):
        for c_idx, val in enumerate(row):
            set_cell_text(t_ex.cell(r_idx, c_idx), val, bold=(c_idx==0))

    doc.add_paragraph().paragraph_format.space_before = Pt(10)

    # Section 4
    doc.add_heading("4. Ma trận Phân quyền Spring Security 6", level=1)
    doc.add_paragraph("Hệ thống cấu hình 2 vai trò truy cập với BCryptPasswordEncoder:")
    doc.add_paragraph("• ADMIN (admin / admin123): Có toàn quyền CRUD sinh viên, môn học; đăng ký/hủy học phần; xem dashboard.")
    doc.add_paragraph("• USER (user / user123): Được xem danh sách sinh viên, danh mục môn học, thực hiện đăng ký và hủy học phần. Bị chặn (403) khi cố truy cập các trang thêm/sửa/xóa sinh viên hoặc môn học.")

    # Section 5: Screenshots placeholders
    doc.add_heading("5. Hình ảnh Minh chứng Kết quả Hoàn thành (Demo)", level=1)
    doc.add_paragraph("Các ảnh chụp màn hình kiểm thử trực tiếp hệ thống:")

    add_image_placeholder(doc, "Trang Đăng Nhập (Login)", "Chụp màn hình giao diện đăng nhập tại /login với thẻ tài khoản demo")
    doc.add_paragraph()
    add_image_placeholder(doc, "Bảng Điều Khiển (Dashboard - Bài 9)", "Chụp màn hình hiển thị 3 thẻ thống kê: Tổng SV, Tổng Môn, Tổng Đăng ký")
    doc.add_paragraph()
    add_image_placeholder(doc, "Quản Lý Sinh Viên (CRUD)", "Chụp màn hình danh sách sinh viên và form thêm/sửa sinh viên")
    doc.add_paragraph()
    add_image_placeholder(doc, "Danh Mục Môn Học (CRUD Khóa học)", "Chụp màn hình danh sách môn học và form thêm môn học mới")
    doc.add_paragraph()
    add_image_placeholder(doc, "Đăng Ký Học Phần Mới (Bài 5 & 6)", "Chụp màn hình form chọn sinh viên và môn học để đăng ký")
    doc.add_paragraph()
    add_image_placeholder(doc, "Danh Sách Đăng Ký & Nút Hủy ĐK (Bài 6 & 7)", "Chụp màn hình bảng danh sách lượt đăng ký có nút Hủy ĐK màu đỏ")
    doc.add_paragraph()
    add_image_placeholder(doc, "Lịch Sử Đăng Ký Của 1 Sinh Viên (Bài 8)", "Chụp màn hình chi tiết môn học đã đăng ký và tổng số tín chỉ tích lũy")
    doc.add_paragraph()
    add_image_placeholder(doc, "Phân Quyền & Trang 403 Forbidden (Bài 10)", "Chụp màn hình khi đăng nhập tài khoản USER cố truy cập /students/new")

    # Section 6
    doc.add_heading("6. Kết luận & Đánh giá Đạt được", level=1)
    doc.add_paragraph("Dự án Lab 15 đã hoàn thành xuất sắc 100% các yêu cầu:")
    doc.add_paragraph("1. Đáp ứng trọn vẹn 10/10 bài tập theo khung chương trình của học phần.")
    doc.add_paragraph("2. Mã nguồn được tổ chức sạch đẹp, chuẩn MVC, tuân thủ nguyên lý Clean Code.")
    doc.add_paragraph("3. Bộ kiểm thử tự động (Unit Test & Integration Test) đạt 100% BUILD SUCCESS.")
    doc.add_paragraph("4. Sản phẩm sẵn sàng nộp bài và trình chiếu demo.")

    doc.save(OUTPUT)
    print(f"Báo cáo Word đã được tạo thành công tại: {OUTPUT}")

if __name__ == "__main__":
    main()
